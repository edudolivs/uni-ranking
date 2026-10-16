"""Prepara os dados agregados que alimentam o dashboard.

Entradas (em dados/):
  arwu/arwu_<ano>.json            API pública do ShanghaiRanking
  igc/*.xlsx                      Índice Geral de Cursos (Inep)
  enade/conceito_enade_<ano>.xlsx Conceito Enade por curso (Inep)
  cpc/*.xlsx                      Conceito Preliminar de Curso (Inep)
  processado/leiden_br_<ed>.csv   gerado por extrair_leiden.py
  processado/qs_br.csv            gerado por extrair_qs.py
  processado/scimago_br.csv       gerado por extrair_scimago.py
  processado/ruf.csv              gerado por extrair_ruf.py
Saída: dados/processado/dados_dashboard.json

Nenhum valor é inventado: o que não existir na fonte fica nulo.
"""
import csv
import glob
import json
import math
import os
import re
import unicodedata
from collections import Counter, defaultdict

import openpyxl

RAIZ = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "dados")
SAIDA = os.path.join(RAIZ, "processado", "dados_dashboard.json")

GRUPOS = {"internacional": ["arwu", "leiden", "qs", "scimago"], "brasileiro": ["igc", "cpc", "enade", "ruf"]}


def num(x):
    """Converte para float; devolve None para vazio, '-', 'SC' etc."""
    if x is None:
        return None
    try:
        return float(str(x).replace(",", "."))
    except ValueError:
        return None


def ano_valido(x):
    """Ano inteiro entre 2000 e 2100; None para notas de rodapé e linhas vazias."""
    try:
        v = int(float(str(x).strip()))
    except (TypeError, ValueError):
        return None
    return v if 2000 <= v <= 2100 else None


def faixa(x):
    if x is None:
        return "SC"
    s = str(x).strip()
    if s in ("", "-"):
        return "SC"
    try:
        return str(int(float(s)))
    except ValueError:
        return "SC"


# ---------------------------------------------------------------- chave de nome
STOP = {
    "university", "universidade", "universidad", "of", "de", "do", "da", "dos", "das", "e", "the",
    "fundacao", "foundation", "instituto", "institute", "centro", "center", "centre", "faculdade",
}
TRAD = {"state": "estadual", "catholic": "catolica", "pontifical": "pontificia", "pontificial": "pontificia",
        "federal": "federal", "estadual": "estadual", "catolica": "catolica", "pontificia": "pontificia"}


# Mesma instituição com nomes sem equivalência por tokens (inglês x português)
ALIAS = {"estadual paulo sao": "estadual paulista"}  # Sao Paulo State University = Unesp


def titulo(nome):
    """Caixa de título em português: 'UNIVERSIDADE FEDERAL DA PARAÍBA' -> 'Universidade Federal da Paraíba'."""
    miudas = {"de", "da", "do", "das", "dos", "e"}
    palavras = str(nome).lower().split()
    return " ".join(w if (i and w in miudas) else w.capitalize() for i, w in enumerate(palavras))


def chave(nome):
    s = unicodedata.normalize("NFKD", str(nome)).encode("ascii", "ignore").decode().lower()
    toks = [TRAD.get(t, t) for t in re.findall(r"[a-z]+", s)]
    toks = [t for t in toks if t not in STOP]
    k = " ".join(sorted(set(toks)))
    return ALIAS.get(k, k)


# ---------------------------------------------------------------- ARWU
def ler_arwu():
    anos, br = {}, []
    for f in sorted(glob.glob(os.path.join(RAIZ, "arwu", "arwu_*.json"))):
        ano = int(re.search(r"(\d{4})", os.path.basename(f)).group(1))
        lista = json.load(open(f, encoding="utf-8"))["data"]["rankings"]
        brasil = [r for r in lista if r.get("region") == "Brazil"]
        paises = Counter(r.get("region") for r in lista)
        anos[ano] = {"total": len(lista), "brasil": len(brasil), "paises": paises.most_common(8)}
        for r in brasil:
            rk = str(r["ranking"])
            br.append({
                "ano": ano,
                "nome": r["univNameEn"],
                "ranking": rk,
                "ranking_min": int(re.match(r"\d+", rk).group()),
                "regional": r.get("regionRanking"),
                "score": num(r.get("score")),
                "ind": r.get("indData"),
            })
    return anos, br


# ---------------------------------------------------------------- IGC
def ler_igc():
    por_ano, regs = {}, []
    arqs = sorted(glob.glob(os.path.join(RAIZ, "igc", "*.xlsx")))
    for f in arqs:
        wb = openpyxl.load_workbook(f, read_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        cab = [str(c).strip().replace("*", "").strip() if c else "" for c in next(it)]
        ix = {c: i for i, c in enumerate(cab)}
        cont = Counter()
        ano = None
        for r in it:
            a = ano_valido(r[ix["Ano"]])
            if a is None:
                continue
            ano = a
            fx = faixa(r[ix["IGC (Faixa)"]])
            cat = r[ix["Categoria Administrativa"]] or "n/d"
            org = r[ix["Organização Acadêmica"]] or "n/d"
            cont[(fx, cat)] += 1
            regs.append({
                "ano": ano, "nome": str(r[ix["Nome da IES"]]).strip(), "sigla": r[ix["Sigla da IES"]],
                "uf": r[ix["Sigla da UF"]], "org": org, "cat": cat,
                "igc": num(r[ix["IGC (Contínuo)"]]), "faixa": fx,
                "cursos_cpc": num(r[ix["Nº de Cursos com CPC no triênio"]]),
                "conceito_grad": num(r[ix["Conceito Médio de Graduação"]]),
                "conceito_mest": num(r[ix["Conceito Médio de Mestrado"]]),
                "conceito_dout": num(r[ix["Conceito Médio do doutorado"]]),
            })
        por_ano[ano] = [{"faixa": k[0], "cat": k[1], "n": v} for k, v in sorted(cont.items())]
    return por_ano, regs


# ---------------------------------------------------------------- ENADE
def ler_enade():
    resumo = defaultdict(lambda: {"cursos": 0, "ies": set(), "soma": 0.0, "n": 0})
    por_ies = defaultdict(lambda: {"cursos": 0, "soma": 0.0, "n": 0, "nome": None, "cat": None, "org": None})
    arqs = sorted(glob.glob(os.path.join(RAIZ, "enade", "conceito_enade_*.xlsx")))
    for f in arqs:
        wb = openpyxl.load_workbook(f, read_only=True)
        aba = next(n for n in wb.sheetnames if n.upper().startswith("PLANILHA") and n != "Planilha2")
        it = wb[aba].iter_rows(values_only=True)
        cab = [str(c).strip().replace("*", "").strip() if c else "" for c in next(it)]
        ix = {c: i for i, c in enumerate(cab)}
        for r in it:
            ano = ano_valido(r[ix["Ano"]])
            if ano is None:
                continue
            cat = r[ix["Categoria Administrativa"]] or "n/d"
            cod = r[ix["Código da IES"]]
            v = num(r[ix["Conceito Enade (Contínuo)"]])
            g = resumo[(ano, cat)]
            g["cursos"] += 1
            g["ies"].add(cod)
            if v is not None:
                g["soma"] += v
                g["n"] += 1
            p = por_ies[(ano, cod)]
            p["nome"], p["cat"], p["org"] = str(r[ix["Nome da IES"]]).strip(), cat, r[ix["Organização Acadêmica"]]
            p["cursos"] += 1
            if v is not None:
                p["soma"] += v
                p["n"] += 1
    res = [{"ano": a, "cat": c, "cursos": g["cursos"], "ies": len(g["ies"]),
            "media": round(g["soma"] / g["n"], 3) if g["n"] else None}
           for (a, c), g in sorted(resumo.items())]
    ies = [{"ano": a, "cod": c, "nome": p["nome"], "cat": p["cat"], "org": p["org"], "cursos": p["cursos"],
            "media": round(p["soma"] / p["n"], 3) if p["n"] else None}
           for (a, c), p in sorted(por_ies.items(), key=lambda kv: (kv[0][0], str(kv[0][1])))]
    return res, ies


# ---------------------------------------------------------------- CPC (Inep)
def ler_cpc():
    """Mesma lógica de ler_enade(), trocando o conceito pelo CPC (Contínuo)."""
    resumo = defaultdict(lambda: {"cursos": 0, "ies": set(), "soma": 0.0, "n": 0})
    por_ies = defaultdict(lambda: {"cursos": 0, "soma": 0.0, "n": 0, "nome": None, "cat": None, "org": None})
    arqs = sorted(glob.glob(os.path.join(RAIZ, "cpc", "*.xlsx")))
    for f in arqs:
        wb = openpyxl.load_workbook(f, read_only=True)
        aba = next(n for n in wb.sheetnames if "cpc" in n.lower())
        it = wb[aba].iter_rows(values_only=True)
        cab = [str(c).strip().replace("*", "").strip() if c else "" for c in next(it)]
        ix = {c: i for i, c in enumerate(cab)}
        for r in it:
            ano = ano_valido(r[ix["Ano"]])
            if ano is None:
                continue
            cat = r[ix["Categoria Administrativa"]] or "n/d"
            cod = r[ix["Código da IES"]]
            v = num(r[ix["CPC (Contínuo)"]])
            g = resumo[(ano, cat)]
            g["cursos"] += 1
            g["ies"].add(cod)
            if v is not None:
                g["soma"] += v
                g["n"] += 1
            p = por_ies[(ano, cod)]
            p["nome"], p["cat"], p["org"] = str(r[ix["Nome da IES"]]).strip(), cat, r[ix["Organização Acadêmica"]]
            p["cursos"] += 1
            if v is not None:
                p["soma"] += v
                p["n"] += 1
    res = [{"ano": a, "cat": c, "cursos": g["cursos"], "ies": len(g["ies"]),
            "media": round(g["soma"] / g["n"], 3) if g["n"] else None}
           for (a, c), g in sorted(resumo.items())]
    ies = [{"ano": a, "cod": c, "nome": p["nome"], "cat": p["cat"], "org": p["org"], "cursos": p["cursos"],
            "media": round(p["soma"] / p["n"], 3) if p["n"] else None}
           for (a, c), p in sorted(por_ies.items(), key=lambda kv: (kv[0][0], str(kv[0][1])))]
    return res, ies


# ---------------------------------------------------------------- Leiden
def ler_leiden():
    regs = []
    for f in sorted(glob.glob(os.path.join(RAIZ, "processado", "leiden_br_*.csv"))):
        for r in csv.DictReader(open(f, encoding="utf-8")):
            regs.append({
                "edicao": int(r["Edicao"]), "nome": r["University"], "periodo": r["Period"],
                "P": num(r.get("P")), "P_top10": num(r.get("P_top10")), "PP_top10": num(r.get("PP_top10")),
                "MNCS": num(r.get("MNCS")), "PP_int_collab": num(r.get("PP_int_collab")),
                "PP_industry_collab": num(r.get("PP_industry_collab")), "PP_OA": num(r.get("PP_OA")),
            })
    return regs


# ---------------------------------------------------------------- QS
def ler_qs():
    caminho = os.path.join(RAIZ, "processado", "qs_br.csv")
    if not os.path.exists(caminho):
        return []
    regs = []
    for r in csv.DictReader(open(caminho, encoding="utf-8")):
        regs.append({
            "ano": int(r["ano"]), "nome": r["nome"],
            "rank_raw": r["rank_raw"], "rank_min": num(r["rank_min"]), "score_geral": num(r["score_geral"]),
            "ar_score": num(r["ar_score"]), "er_score": num(r["er_score"]),
            "fsr_score": num(r["fsr_score"]), "cpf_score": num(r["cpf_score"]),
        })
    return regs


# ---------------------------------------------------------------- SCImago
def ler_scimago():
    caminho = os.path.join(RAIZ, "processado", "scimago_br.csv")
    if not os.path.exists(caminho):
        return []
    regs = []
    for r in csv.DictReader(open(caminho, encoding="utf-8")):
        regs.append({
            "ano": int(r["ano"]), "nome": r["nome"],
            "rank_br": num(r["rank_br"]), "rank_global": num(r["rank_global"]), "quartil": r["quartil"] or None,
        })
    return regs


# ---------------------------------------------------------------- RUF
def ler_ruf():
    caminho = os.path.join(RAIZ, "processado", "ruf.csv")
    if not os.path.exists(caminho):
        return []
    regs = []
    for r in csv.DictReader(open(caminho, encoding="utf-8")):
        regs.append({k: (int(r["ano"]) if k == "ano" else (r[k] if k in ("nome", "categoria") else num(r[k])))
                      for k in r})
    return regs


# ---------------------------------------------------------------- cruzamento (N fontes)
def cruzar(fontes_mais_recentes):
    """fontes_mais_recentes: {nome_fonte: {chave: {"valor": ..., "nome": ...}}} já restrito à edição/ano mais recente.
    Une por chave de nome as universidades brasileiras presentes em 2+ fontes."""
    linhas = defaultdict(dict)
    nomes = {}
    for fonte, registros in fontes_mais_recentes.items():
        for k, d in registros.items():
            linhas[k][fonte] = d["valor"]
            if fonte in ("igc", "cpc", "enade"):
                nomes[k] = titulo(d["nome"])
            else:
                nomes.setdefault(k, d["nome"])
    saida = []
    for k, d in linhas.items():
        if len(d) < 2:
            continue
        saida.append({"nome": nomes[k], "fontes": sorted(d.keys()), **{f"v_{f}": v for f, v in d.items()}})
    saida.sort(key=lambda r: -len(r["fontes"]))
    cobertura = {f: sum(1 for d in linhas.values() if f in d) for f in fontes_mais_recentes}
    cobertura["em_2_ou_mais"] = len(saida)
    cobertura["em_todas"] = sum(1 for r in saida if len(r["fontes"]) == len(fontes_mais_recentes))
    return saida, cobertura


def mais_recente(registros, campo_ano, campo_nome, campo_valor, filtro=None):
    """{chave: {"valor":..., "nome":...}} usando só o maior campo_ano (e, se dado, só linhas em que filtro(r) é True)."""
    validos = [r for r in registros if r.get(campo_valor) is not None and (filtro is None or filtro(r))]
    if not validos:
        return {}, None
    ult = max(r[campo_ano] for r in validos)
    out = {}
    for r in validos:
        if r[campo_ano] != ult:
            continue
        out[chave(r[campo_nome])] = {"valor": r[campo_valor], "nome": r[campo_nome]}
    return out, ult


# ---------------------------------------------------------------- correlação de Spearman
def _postos(valores):
    """Rank médio (1-based) com empate resolvido pela média das posições empatadas."""
    ordem = sorted(range(len(valores)), key=lambda i: valores[i])
    postos = [0.0] * len(valores)
    i = 0
    while i < len(ordem):
        j = i
        while j + 1 < len(ordem) and valores[ordem[j + 1]] == valores[ordem[i]]:
            j += 1
        media = (i + j) / 2 + 1
        for k in range(i, j + 1):
            postos[ordem[k]] = media
        i = j + 1
    return postos


def spearman(pares):
    """pares: lista de (x, y). Retorna (rho, n); (None, n) se n < 5 ou variância zero."""
    pares = [(x, y) for x, y in pares if x is not None and y is not None]
    n = len(pares)
    if n < 5:
        return None, n
    xs, ys = _postos([p[0] for p in pares]), _postos([p[1] for p in pares])
    mx, my = sum(xs) / n, sum(ys) / n
    cov = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    vx = sum((x - mx) ** 2 for x in xs)
    vy = sum((y - my) ** 2 for y in ys)
    if vx == 0 or vy == 0:
        return None, n
    return cov / math.sqrt(vx * vy), n


def serie_por_ano(registros, campo_ano, campo_nome, campo_valor, filtro=None):
    """{ano: {chave: valor}}"""
    s = {}
    for r in registros:
        if r.get(campo_valor) is None or (filtro and not filtro(r)):
            continue
        s.setdefault(r[campo_ano], {})[chave(r[campo_nome])] = r[campo_valor]
    return s


def serie_leiden_ultimo_periodo(leiden, campo):
    por_ed = defaultdict(set)
    for r in leiden:
        por_ed[r["edicao"]].add(r["periodo"])
    ult_periodo = {ed: max(ps) for ed, ps in por_ed.items()}
    s = {}
    for r in leiden:
        if r["periodo"] != ult_periodo[r["edicao"]] or r.get(campo) is None:
            continue
        s.setdefault(r["edicao"], {})[chave(r["nome"])] = r[campo]
    return s


def correl_estabilidade(serie, rotulo):
    """Spearman da mesma métrica entre edições consecutivas (estabilidade temporal, hipótese c)."""
    out = []
    anos = sorted(serie)
    for a0, a1 in zip(anos, anos[1:]):
        comuns = set(serie[a0]) & set(serie[a1])
        rho, n = spearman([(serie[a0][k], serie[a1][k]) for k in comuns])
        if rho is not None:
            out.append({"rotulo": rotulo, "de": a0, "para": a1, "rho": round(rho, 3), "n": n})
    return out


def correl_par(serie_a, serie_b, rotulo, usar_ano_de="a"):
    """Spearman entre duas séries {ano: {chave: valor}} no(s) ano(s) comuns a ambas."""
    out = []
    for ano in sorted(set(serie_a) & set(serie_b)):
        comuns = set(serie_a[ano]) & set(serie_b[ano])
        rho, n = spearman([(serie_a[ano][k], serie_b[ano][k]) for k in comuns])
        if rho is not None:
            out.append({"rotulo": rotulo, "ano": ano, "rho": round(rho, 3), "n": n})
    return out


def correlacionar(arwu_br, leiden, qs, scimago, igc, cpc_ies, enade_ies, ruf):
    """Spearman com o que está disponível hoje. 'posição'/'ranking' usa número menor = melhor;
    'nota'/'score'/indicador usa número maior = melhor -- por isso o sinal de rho indica a direção
    (positivo = duas posições concordam; negativo = indicador alto associado a posição numericamente
    menor, ou seja, melhor). n é sempre reportado ao lado: amostras pequenas não geram p-value."""
    s_arwu = serie_por_ano(arwu_br, "ano", "nome", "ranking_min")
    s_leiden_pp = serie_leiden_ultimo_periodo(leiden, "PP_top10")
    s_leiden_p = serie_leiden_ultimo_periodo(leiden, "P")
    s_qs_rank = serie_por_ano(qs, "ano", "nome", "rank_min")
    s_qs_cpf = serie_por_ano(qs, "ano", "nome", "cpf_score")
    s_qs_ar = serie_por_ano(qs, "ano", "nome", "ar_score")
    s_scimago = serie_por_ano(scimago, "ano", "nome", "rank_global")
    s_igc = serie_por_ano(igc, "ano", "nome", "igc", filtro=lambda r: r["org"] == "Universidade")
    s_cpc = serie_por_ano(cpc_ies, "ano", "nome", "media", filtro=lambda r: r["org"] == "Universidade")
    s_enade = serie_por_ano(enade_ies, "ano", "nome", "media", filtro=lambda r: r["org"] == "Universidade")
    s_ruf_rank = serie_por_ano(ruf, "ano", "nome", "ranking")
    s_ruf_ensino = serie_por_ano(ruf, "ano", "nome", "nota_ensino")
    s_ruf_pesquisa = serie_por_ano(ruf, "ano", "nome", "nota_pesquisa")
    s_ruf_mercado = serie_por_ano(ruf, "ano", "nome", "nota_mercado")

    resultado = {"estabilidade": [], "a_pesquisa_internacional": [], "b_ensino_mercado_brasileiro": [], "entre_rankings": []}

    for serie, rot in [(s_arwu, "ARWU: posição ano a ano"), (s_qs_rank, "QS: posição ano a ano"),
                        (s_scimago, "SCImago: posição global ano a ano"), (s_ruf_rank, "RUF: posição ano a ano")]:
        resultado["estabilidade"] += correl_estabilidade(serie, rot)

    resultado["a_pesquisa_internacional"] += correl_par(s_qs_rank, s_qs_cpf, "QS: posição × citações/docente")
    resultado["a_pesquisa_internacional"] += correl_par(s_qs_rank, s_qs_ar, "QS: posição × reputação acadêmica")
    resultado["a_pesquisa_internacional"] += correl_par(s_arwu, s_leiden_pp, "ARWU × Leiden: posição × % top 10%")
    resultado["a_pesquisa_internacional"] += correl_par(s_arwu, s_leiden_p, "ARWU × Leiden: posição × publicações")

    resultado["b_ensino_mercado_brasileiro"] += correl_par(s_ruf_rank, s_ruf_ensino, "RUF: posição × nota de ensino")
    resultado["b_ensino_mercado_brasileiro"] += correl_par(s_ruf_rank, s_ruf_mercado, "RUF: posição × nota de mercado")
    resultado["b_ensino_mercado_brasileiro"] += correl_par(s_ruf_rank, s_ruf_pesquisa, "RUF: posição × nota de pesquisa")

    resultado["entre_rankings"] += correl_par(s_arwu, s_qs_rank, "ARWU × QS (posição, mesma edição)")
    resultado["entre_rankings"] += correl_par(s_arwu, s_scimago, "ARWU × SCImago (posição, mesma edição)")
    resultado["entre_rankings"] += correl_par(s_qs_rank, s_scimago, "QS × SCImago (posição, mesma edição)")
    resultado["entre_rankings"] += correl_par(s_igc, s_ruf_rank, "IGC × RUF (nota IGC × posição RUF)")
    resultado["entre_rankings"] += correl_par(s_igc, s_cpc, "IGC × CPC (notas, mesmo ano)")
    resultado["entre_rankings"] += correl_par(s_igc, s_enade, "IGC × Enade (notas, mesmo ano)")

    return resultado


def main():
    arwu_anos, arwu_br = ler_arwu()
    igc_ano, igc = ler_igc()
    enade_res, enade_ies = ler_enade()
    cpc_res, cpc_ies = ler_cpc()
    leiden = ler_leiden()
    qs = ler_qs()
    scimago = ler_scimago()
    ruf = ler_ruf()

    leiden_mais_recente_edicao = max((r["edicao"] for r in leiden), default=None)
    leiden_ultimo = [r for r in leiden if r["edicao"] == leiden_mais_recente_edicao]
    leiden_ultimo_periodo = max((r["periodo"] for r in leiden_ultimo), default=None)

    fontes_mais_recentes = {}
    fontes_mais_recentes["arwu"], _ = mais_recente(arwu_br, "ano", "nome", "ranking_min")
    fontes_mais_recentes["leiden"], _ = mais_recente(
        [r for r in leiden if r["edicao"] == leiden_mais_recente_edicao and r["periodo"] == leiden_ultimo_periodo],
        "edicao", "nome", "PP_top10")
    fontes_mais_recentes["igc"], _ = mais_recente(igc, "ano", "nome", "igc", filtro=lambda r: r["org"] == "Universidade")
    fontes_mais_recentes["cpc"], _ = mais_recente(cpc_ies, "ano", "nome", "media", filtro=lambda r: r["org"] == "Universidade")
    fontes_mais_recentes["enade"], _ = mais_recente(enade_ies, "ano", "nome", "media", filtro=lambda r: r["org"] == "Universidade")
    fontes_mais_recentes["qs"], _ = mais_recente(qs, "ano", "nome", "rank_min")
    fontes_mais_recentes["scimago"], _ = mais_recente(scimago, "ano", "nome", "rank_global")
    fontes_mais_recentes["ruf"], _ = mais_recente(ruf, "ano", "nome", "ranking")
    fontes_mais_recentes = {k: v for k, v in fontes_mais_recentes.items() if v}

    cruz, cob = cruzar(fontes_mais_recentes)
    cob["anos"] = {f: (max(r["ano"] for r in regs) if f != "leiden" else leiden_mais_recente_edicao)
                   for f, regs in [("arwu", arwu_br), ("igc", igc), ("cpc", cpc_ies), ("enade", enade_ies),
                                    ("leiden", leiden), ("qs", qs), ("scimago", scimago), ("ruf", ruf)]
                   if regs and f in fontes_mais_recentes}

    correlacoes = correlacionar(arwu_br, leiden, qs, scimago, igc, cpc_ies, enade_ies, ruf)

    saida = {
        "arwu": {"anos": arwu_anos, "brasil": arwu_br},
        "igc": {"por_ano": igc_ano, "universidades": [r for r in igc if r["org"] == "Universidade"]},
        "enade": {"resumo": enade_res},
        "cpc": {"resumo": cpc_res},
        "leiden": leiden,
        "qs": qs,
        "scimago": scimago,
        "ruf": ruf,
        "cruzamento": cruz,
        "cobertura": cob,
        "grupos": GRUPOS,
        "correlacoes": correlacoes,
    }
    with open(SAIDA, "w", encoding="utf-8") as f:
        json.dump(saida, f, ensure_ascii=False, separators=(",", ":"))
    print("salvo", os.path.abspath(SAIDA), round(os.path.getsize(SAIDA) / 1024), "KB")
    print("ARWU brasileiras por ano:", {a: v["brasil"] for a, v in arwu_anos.items()})
    print("IGC linhas por ano:", {a: sum(x["n"] for x in v) for a, v in igc_ano.items()})
    print("ENADE cursos por ano:", {a: sum(r["cursos"] for r in enade_res if r["ano"] == a) for a in sorted({r['ano'] for r in enade_res})})
    print("CPC cursos por ano:", {a: sum(r["cursos"] for r in cpc_res if r["ano"] == a) for a in sorted({r['ano'] for r in cpc_res})})
    print("Leiden linhas:", len(leiden), "| QS linhas:", len(qs), "| SCImago linhas:", len(scimago), "| RUF linhas:", len(ruf))
    print("cobertura do cruzamento:", cob)
    print("correlações calculadas:", {k: len(v) for k, v in correlacoes.items()})


if __name__ == "__main__":
    main()
